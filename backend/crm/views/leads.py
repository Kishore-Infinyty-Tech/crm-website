"""Leads: list, create, detail, edit, stage changes (with handover), contacts, import."""
import csv
import io
import re

from django.conf import settings

from ..common import (
    ApiError,
    build_update,
    clean,
    get_lead,
    lead_scope,
    require,
    to_int,
    to_num,
)
from ..db import execute, fetch_all, fetch_one, transaction
from ..emails import send_async
from ..http import api, ok

LIST_COLUMNS = """
    l.id, l.project_summary, l.estimated_budget, l.priority, l.stage_id,
    ps.name AS stage, ps.position AS stage_position, l.stage_changed_at, l.created_at,
    l.assigned_to, au.full_name AS assigned_name, l.created_by,
    c.id AS company_id, c.name AS company, c.industry, c.location,
    ct.name AS contact_name, ct.phone, ct.email,
    src.name AS source, l.source_id,
    (SELECT min(f.due_at) FROM follow_ups f
      WHERE f.lead_id = l.id AND f.status = 'pending') AS next_follow_up
"""
LIST_FROM = """
    FROM leads l
    JOIN companies c ON c.id = l.company_id
    JOIN pipeline_stages ps ON ps.id = l.stage_id
    LEFT JOIN contacts ct ON ct.id = l.primary_contact_id
    LEFT JOIN lead_sources src ON src.id = l.source_id
    LEFT JOIN users au ON au.id = l.assigned_to
"""
LIST_SQL = f"SELECT {LIST_COLUMNS} {LIST_FROM}"

# Lead page: the lead, its people and its stage history in one trip to the database.
DETAIL_SQL = f"""
    SELECT {LIST_COLUMNS},
      c.website, l.deal_value, l.closed_at, l.lost_reason_id, l.primary_contact_id,
      (SELECT coalesce(json_agg(x ORDER BY x.is_primary DESC, x.id), '[]'::json) FROM (
         SELECT id, name, job_title, email, phone, alt_phone, is_primary
         FROM contacts WHERE company_id = l.company_id) x) AS contacts,
      (SELECT coalesce(json_agg(h ORDER BY h.changed_at DESC, h.id DESC), '[]'::json) FROM (
         SELECT h.id, fs.name AS from_stage, ts.name AS to_stage, u.full_name AS changed_by, h.changed_at, h.note
         FROM lead_stage_history h
         LEFT JOIN pipeline_stages fs ON fs.id = h.from_stage
         JOIN pipeline_stages ts ON ts.id = h.to_stage
         LEFT JOIN users u ON u.id = h.changed_by
         WHERE h.lead_id = l.id) h) AS history
    {LIST_FROM}
"""


def digits(phone):
    return re.sub(r"[^0-9]", "", phone or "")


def find_duplicate(conn, email, phone):
    """Return an existing lead that shares this email or phone, or None."""
    conditions, params = [], []
    if email:
        conditions.append("lower(ct.email) = lower(%s)")
        params.append(email)
    if len(digits(phone)) >= 7:
        # Compare the last 10 digits so "+91 98765 00011" matches "9876500011".
        conditions.append("right(regexp_replace(ct.phone, '[^0-9]', '', 'g'), 10) = %s")
        params.append(digits(phone)[-10:])
    if not conditions:
        return None
    return fetch_one(
        "SELECT l.id AS lead_id, c.name AS company FROM contacts ct "
        "JOIN companies c ON c.id = ct.company_id JOIN leads l ON l.company_id = c.id "
        f"WHERE {' OR '.join(conditions)} LIMIT 1",
        params,
        conn=conn,
    )


def create_lead_records(conn, user, data, *, source_id=None, import_batch_id=None):
    """Insert company + primary contact + lead + first stage-history row. Returns lead id."""
    company = execute(
        "INSERT INTO companies (name, website, industry, location, created_by) "
        "VALUES (%s, %s, %s, %s, %s) RETURNING id",
        [
            data["company_name"].strip(),
            clean(data.get("website")),
            clean(data.get("industry")),
            clean(data.get("location")),
            user["id"],
        ],
        conn=conn,
    )
    contact_id = None
    if clean(data.get("contact_name")):
        contact = execute(
            "INSERT INTO contacts (company_id, name, job_title, email, phone, alt_phone, is_primary) "
            "VALUES (%s, %s, %s, %s, %s, %s, TRUE) RETURNING id",
            [
                company["id"],
                data["contact_name"].strip(),
                clean(data.get("job_title")),
                clean(data.get("email")),
                clean(data.get("phone")),
                clean(data.get("alt_phone")),
            ],
            conn=conn,
        )
        contact_id = contact["id"]
    stage = fetch_one("SELECT id FROM pipeline_stages WHERE position = 1", conn=conn)
    assigned = to_int(data.get("assigned_to"), "assigned to") or user["id"]
    lead = execute(
        "INSERT INTO leads (company_id, primary_contact_id, source_id, stage_id, project_summary, "
        "estimated_budget, priority, created_by, assigned_to, import_batch_id) "
        "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING id",
        [
            company["id"],
            contact_id,
            source_id or to_int(data.get("source_id"), "source"),
            stage["id"],
            clean(data.get("project_summary")),
            to_num(data.get("estimated_budget"), "budget"),
            clean(data.get("priority")),
            user["id"],
            assigned,
            import_batch_id,
        ],
        conn=conn,
    )
    execute(
        "INSERT INTO lead_stage_history (lead_id, from_stage, to_stage, changed_by, note) "
        "VALUES (%s, NULL, %s, %s, 'Lead created')",
        [lead["id"], stage["id"], user["id"]],
        conn=conn,
    )
    return lead["id"]


# ---------------- list / create ----------------
@api("GET", "POST")
def leads(request):
    user = request.user_row
    if request.method == "POST":
        return create_lead(request)

    scope_sql, params = lead_scope(user)
    where = [scope_sql]
    get = request.GET
    for param, column in (("stage_id", "l.stage_id"), ("assigned_to", "l.assigned_to"), ("source_id", "l.source_id")):
        if get.get(param):
            where.append(f"{column} = %s")
            params.append(to_int(get[param], param))
    if get.get("priority"):
        where.append("l.priority = %s")
        params.append(get["priority"])
    if get.get("mine") == "1":
        where.append("l.assigned_to = %s")
        params.append(user["id"])
    if get.get("q"):
        pattern = f"%{get['q'].strip()}%"
        where.append("(c.name ILIKE %s OR ct.name ILIKE %s OR ct.phone ILIKE %s OR ct.email ILIKE %s)")
        params += [pattern] * 4
    rows = fetch_all(f"{LIST_SQL} WHERE {' AND '.join(where)} ORDER BY l.created_at DESC LIMIT 1000", params)
    return ok({"leads": rows})


def create_lead(request):
    user = request.user_row
    data = request.json
    require(data, "company_name")
    if not clean(data.get("phone")) and not clean(data.get("email")):
        raise ApiError("Add at least a phone number or an email so the team can contact them")
    with transaction() as conn:
        if not data.get("force"):
            dup = find_duplicate(conn, clean(data.get("email")), clean(data.get("phone")))
            if dup:
                raise ApiError(
                    f"A lead with this phone or email already exists ({dup['company']})",
                    409,
                    duplicate_lead_id=dup["lead_id"],
                )
        lead_id = create_lead_records(conn, user, data)
    return ok({"id": lead_id}, 201)


# ---------------- detail ----------------
@api("GET", "PATCH", "DELETE")
def lead_detail(request, lead_id):
    user = request.user_row
    if request.method == "GET":
        scope_sql, scope_params = lead_scope(user)
        detail = fetch_one(f"{DETAIL_SQL} WHERE l.id = %s AND {scope_sql}", [lead_id, *scope_params])
        if not detail:
            raise ApiError("Lead not found", 404)
        return ok(detail)

    lead = get_lead(user, lead_id)

    if request.method == "DELETE":
        if user["role"] != "admin":
            raise ApiError("Only CRM Administrators can delete leads", 403)
        with transaction() as conn:
            execute("DELETE FROM leads WHERE id = %s", [lead_id], conn=conn)
            execute(
                "DELETE FROM companies WHERE id = %s AND NOT EXISTS "
                "(SELECT 1 FROM leads WHERE company_id = %s)",
                [lead["company_id"], lead["company_id"]],
                conn=conn,
            )
        return ok()

    if request.method == "PATCH":
        data = request.json
        with transaction() as conn:
            sets, params = build_update(
                "leads",
                ("project_summary", "estimated_budget", "priority", "source_id", "primary_contact_id", "deal_value"),
                data,
                convert={"estimated_budget": to_num, "deal_value": to_num, "source_id": to_int, "primary_contact_id": to_int},
            )
            if "assigned_to" in data and to_int(data["assigned_to"], "assigned to") != lead["assigned_to"]:
                if user["role"] == "executive":
                    raise ApiError("Only CRM Administrators and the CRM Analyst can reassign leads", 403)
                sets.append("assigned_to = %s")
                params.append(to_int(data["assigned_to"], "assigned to"))
            if sets:
                execute(f"UPDATE leads SET {', '.join(sets)} WHERE id = %s", [*params, lead_id], conn=conn)
            c_sets, c_params = build_update(
                "companies", ("website", "industry", "location"), data
            )
            if clean(data.get("company_name")):
                c_sets.append("name = %s")
                c_params.append(data["company_name"].strip())
            if c_sets:
                execute(
                    f"UPDATE companies SET {', '.join(c_sets)} WHERE id = %s",
                    [*c_params, lead["company_id"]],
                    conn=conn,
                )
        return ok()

    raise ApiError("Method not allowed", 405)


# ---------------- stage change with handover ----------------
def pick_owner(conn, role_id, prefer_user=None):
    """Least-loaded active user with the role, or the acting user if they have it."""
    if prefer_user and prefer_user["role_id"] == role_id:
        return prefer_user["id"], prefer_user["full_name"], prefer_user["email"]
    row = fetch_one(
        "SELECT u.id, u.full_name, u.email, "
        "(SELECT count(*) FROM leads l JOIN pipeline_stages s ON s.id = l.stage_id "
        " WHERE l.assigned_to = u.id AND s.name NOT IN ('won','lost')) AS open_leads "
        "FROM users u WHERE u.role_id = %s AND u.is_active ORDER BY open_leads, u.id LIMIT 1",
        [role_id],
        conn=conn,
    )
    return (row["id"], row["full_name"], row["email"]) if row else (None, None, None)


@api("POST")
def lead_stage(request, lead_id):
    user = request.user_row
    lead = get_lead(user, lead_id)
    data = request.json
    require(data, "stage_id")
    new_stage_id = to_int(data["stage_id"], "stage")
    notify = None
    with transaction() as conn:
        stage = fetch_one("SELECT * FROM pipeline_stages WHERE id = %s", [new_stage_id], conn=conn)
        if not stage:
            raise ApiError("Unknown stage", 404)
        if stage["id"] == lead["stage_id"]:
            raise ApiError("The lead is already in this stage")
        if stage["name"] == "lost" and not to_int(data.get("lost_reason_id")):
            raise ApiError("Pick a reason for losing this lead")
        closing = stage["name"] in ("won", "lost")
        execute(
            "UPDATE leads SET stage_id = %s, stage_changed_at = now(), "
            "closed_at = CASE WHEN %s THEN now() ELSE NULL END, "
            "lost_reason_id = %s, deal_value = COALESCE(%s, deal_value) WHERE id = %s",
            [
                stage["id"],
                closing,
                to_int(data.get("lost_reason_id")) if stage["name"] == "lost" else None,
                to_num(data.get("deal_value"), "deal value") if stage["name"] == "won" else None,
                lead_id,
            ],
            conn=conn,
        )
        execute(
            "INSERT INTO lead_stage_history (lead_id, from_stage, to_stage, changed_by, note) "
            "VALUES (%s, %s, %s, %s, %s)",
            [lead_id, lead["stage_id"], stage["id"], user["id"], clean(data.get("note"))],
            conn=conn,
        )
        # Handover: if the new stage belongs to another role, give the lead to that role.
        current = fetch_one("SELECT role_id FROM users WHERE id = %s", [lead["assigned_to"]], conn=conn)
        if stage["owner_role"] and not closing and (not current or current["role_id"] != stage["owner_role"]):
            owner_id, owner_name, owner_email = pick_owner(conn, stage["owner_role"], user)
            if owner_id:
                execute("UPDATE leads SET assigned_to = %s WHERE id = %s", [owner_id, lead_id], conn=conn)
                if owner_id != user["id"]:
                    notify = (owner_email, stage["name"])
        company = fetch_one("SELECT name FROM companies WHERE id = %s", [lead["company_id"]], conn=conn)["name"]
    if notify and notify[0]:
        send_async(
            notify[0],
            f"New lead handed over: {company}",
            f"{user['full_name']} moved {company} to '{notify[1]}' and assigned it to you.\n\n"
            f"Open it: {settings.APP_URL}/leads/{lead_id}\n",
            kind="handover",
            lead_id=lead_id,
        )
    return ok()


# ---------------- contacts ----------------
@api("POST")
def lead_contacts(request, lead_id):
    lead = get_lead(request.user_row, lead_id)
    data = request.json
    require(data, "name")
    with transaction() as conn:
        if data.get("is_primary"):
            execute("UPDATE contacts SET is_primary = FALSE WHERE company_id = %s", [lead["company_id"]], conn=conn)
        row = execute(
            "INSERT INTO contacts (company_id, name, job_title, email, phone, alt_phone, is_primary) "
            "VALUES (%s, %s, %s, %s, %s, %s, %s) RETURNING id",
            [
                lead["company_id"],
                data["name"].strip(),
                clean(data.get("job_title")),
                clean(data.get("email")),
                clean(data.get("phone")),
                clean(data.get("alt_phone")),
                bool(data.get("is_primary")),
            ],
            conn=conn,
        )
        if data.get("is_primary"):
            execute("UPDATE leads SET primary_contact_id = %s WHERE id = %s", [row["id"], lead_id], conn=conn)
    return ok({"id": row["id"]}, 201)


@api("PATCH", "DELETE")
def contact_item(request, lead_id, contact_id):
    lead = get_lead(request.user_row, lead_id)
    contact = fetch_one(
        "SELECT id FROM contacts WHERE id = %s AND company_id = %s", [contact_id, lead["company_id"]]
    )
    if not contact:
        raise ApiError("Contact not found", 404)
    with transaction() as conn:
        if request.method == "DELETE":
            execute("DELETE FROM contacts WHERE id = %s", [contact_id], conn=conn)
            return ok()
        data = request.json
        sets, params = build_update("contacts", ("name", "job_title", "email", "phone", "alt_phone"), data)
        if data.get("is_primary"):
            execute("UPDATE contacts SET is_primary = FALSE WHERE company_id = %s", [lead["company_id"]], conn=conn)
            sets.append("is_primary = TRUE")
            execute("UPDATE leads SET primary_contact_id = %s WHERE id = %s", [contact_id, lead_id], conn=conn)
        if sets:
            execute(f"UPDATE contacts SET {', '.join(sets)} WHERE id = %s", [*params, contact_id], conn=conn)
    return ok()


# ---------------- Excel / CSV import ----------------
IMPORT_COLUMNS = [
    "company_name", "contact_name", "job_title", "phone", "email", "website",
    "industry", "location", "project_summary", "estimated_budget", "priority", "source",
]
ALIASES = {
    "company": "company_name", "company name": "company_name", "organisation": "company_name",
    "organization": "company_name", "business": "company_name", "name": "contact_name",
    "contact": "contact_name", "contact name": "contact_name", "person": "contact_name",
    "title": "job_title", "job title": "job_title", "designation": "job_title",
    "mobile": "phone", "phone number": "phone", "mobile number": "phone", "contact number": "phone",
    "email id": "email", "email address": "email", "mail": "email",
    "site": "website", "url": "website", "city": "location", "address": "location",
    "requirement": "project_summary", "requirements": "project_summary", "summary": "project_summary",
    "description": "project_summary", "notes": "project_summary", "budget": "estimated_budget",
    "lead source": "source", "temperature": "priority",
}


def normalize_header(header):
    key = re.sub(r"[\s_]+", " ", str(header or "").strip().lower())
    return ALIASES.get(key, key.replace(" ", "_"))


def read_rows(upload):
    name = upload.name.lower()
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        return [list(row) for row in reader]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        sheet = load_workbook(upload, read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    raise ApiError("Upload a .csv or .xlsx file")


@api("POST")
def import_leads(request):
    user = request.user_row
    upload = request.FILES.get("file")
    if not upload:
        raise ApiError("Choose a file to import")
    rows = [r for r in read_rows(upload) if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise ApiError("The file has no data rows")
    headers = [normalize_header(h) for h in rows[0]]
    if "company_name" not in headers:
        raise ApiError("The file needs a 'Company' column in the first row")

    sources = {s["name"].lower(): s["id"] for s in fetch_all("SELECT id, name FROM lead_sources")}
    default_source = sources.get("excel")
    imported = skipped = 0
    problems = []
    with transaction() as conn:
        batch = execute(
            "INSERT INTO import_batches (file_name, uploaded_by, total_rows) VALUES (%s, %s, %s) RETURNING id",
            [upload.name, user["id"], len(rows) - 1],
            conn=conn,
        )
        for line_no, row in enumerate(rows[1:], start=2):
            data = {h: (row[i].strip() if i < len(row) and isinstance(row[i], str) else (row[i] if i < len(row) else None)) for i, h in enumerate(headers)}
            if not clean(data.get("company_name")):
                problems.append(f"Row {line_no}: company is empty")
                skipped += 1
                continue
            if not clean(data.get("phone")) and not clean(data.get("email")):
                problems.append(f"Row {line_no}: needs a phone or an email")
                skipped += 1
                continue
            dup = find_duplicate(conn, clean(data.get("email")), clean(data.get("phone")))
            if dup:
                problems.append(f"Row {line_no}: already exists as '{dup['company']}'")
                skipped += 1
                continue
            level = (clean(data.get("priority")) or "").lower()
            data["priority"] = {"hot": "high", "warm": "medium", "cold": "low"}.get(level, level)
            if data["priority"] not in ("high", "medium", "low"):
                data["priority"] = None
            try:
                data["estimated_budget"] = to_num(data.get("estimated_budget"))
            except ApiError:
                data["estimated_budget"] = None
            try:
                with conn.transaction():  # a bad row rolls back alone
                    create_lead_records(
                        conn,
                        user,
                        data,
                        source_id=sources.get((clean(data.get("source")) or "").lower(), default_source),
                        import_batch_id=batch["id"],
                    )
            except Exception as exc:  # noqa: BLE001
                problems.append(f"Row {line_no}: could not be saved ({str(exc)[:80]})")
                skipped += 1
                continue
            imported += 1
        execute(
            "UPDATE import_batches SET imported_rows = %s, skipped_rows = %s, error_details = %s WHERE id = %s",
            [imported, skipped, "\n".join(problems[:200]) or None, batch["id"]],
            conn=conn,
        )
    return ok({"imported": imported, "skipped": skipped, "problems": problems[:50]})


@api("GET")
def import_template(request):
    from django.http import HttpResponse

    sample = [
        IMPORT_COLUMNS,
        ["Acme Traders", "Priya Sharma", "Owner", "9876543210", "priya@acme.example", "https://acme.example",
         "Retail", "Bengaluru", "Needs an online ordering website", "150000", "medium", "excel"],
    ]
    buffer = io.StringIO()
    csv.writer(buffer).writerows(sample)
    response = HttpResponse(buffer.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="leads-template.csv"'
    return response
